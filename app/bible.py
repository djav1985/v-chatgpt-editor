# bible.py
import json
import os
import re
import shutil
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from api import communicate_json, communicate_with_openai, get_client
from docx_markup import serialize_paragraph

_H1 = re.compile(r"^<h1>(.*)</h1>$")
_CHAPTER_WORD = re.compile(r"\bchapter\b", re.IGNORECASE)
_TAGS = re.compile(r"<[^>]+>")
_OFF_VALUES = {"0", "false", "off"}

_STRING_LIST = {"type": "array", "items": {"type": "string"}}


def _record(**fields: dict[str, Any]) -> dict[str, Any]:
    return {
        "type": "object",
        "properties": fields,
        "required": list(fields),
        "additionalProperties": False,
    }


_STR = {"type": "string"}

_CHARACTER_CATEGORIES = [
    "appearance",
    "age",
    "occupation",
    "relationship",
    "background",
    "injury",
    "possession",
    "other",
]
_SETTING_CATEGORIES = ["layout", "security", "location", "description", "other"]


def _named_facts(name_key: str, categories: list[str]) -> dict[str, Any]:
    fact = _record(
        category={"type": "string", "enum": categories},
        fact=_STR,
        quote=_STR,
    )
    return _record(
        **{name_key: _STR},
        aliases=_STRING_LIST,
        facts={"type": "array", "items": fact},
    )


CHAPTER_SCHEMA: dict[str, Any] = {
    "type": "object",
    "properties": {
        "summary": {"type": "string"},
        "characters": {
            "type": "array",
            "items": _named_facts("name", _CHARACTER_CATEGORIES),
        },
        "settings": {
            "type": "array",
            "items": _named_facts("name", _SETTING_CATEGORIES),
        },
        "timeline": {
            "type": "array",
            "items": _record(event=_STR, when=_STR, quote=_STR),
        },
        "knowledge": {
            "type": "array",
            "items": _record(character=_STR, knows=_STR, quote=_STR),
        },
        "threads": {
            "type": "array",
            "items": _record(
                thread=_STR,
                status={
                    "type": "string",
                    "enum": ["opened", "advanced", "resolved"],
                },
                quote=_STR,
            ),
        },
        "style": _record(
            tense=_STR, point_of_view=_STR, spelling=_STR, conventions=_STRING_LIST
        ),
    },
    "required": [
        "summary",
        "characters",
        "settings",
        "timeline",
        "knowledge",
        "threads",
        "style",
    ],
    "additionalProperties": False,
}

CHAPTER_SYSTEM = (
    "You are a continuity editor building a story bible from one chapter of a "
    "manuscript. Record only facts stated in the text; never infer or "
    "guess. Every fact, event, piece of knowledge, and thread needs a "
    "'quote': a short passage copied word for word from the chapter, with "
    "no paraphrase and no markup tags. Characters: the canonical name "
    "exactly as spelled in the text, every alias or spelling variant used "
    "in this chapter, and one fact per item, categorized, covering "
    "appearance (hair, eye color, build), age, occupation, relationships, "
    "background and history, injuries, and possessions. Settings: places "
    "and their fixed details such as layout, security measures, and "
    "location. Timeline: events with the stated day, time, season, or "
    "elapsed time in 'when' (use an empty string if none is stated). "
    "Knowledge: what a named character learns, knows, or does not know "
    "yet. Threads: mysteries, promises, and setups that are opened, "
    "advanced, or resolved. Style: tense, point of view, spelling variant, "
    "and conventions for dialogue and titles. The summary is two or three "
    "neutral sentences and is not used as evidence. Use empty lists when "
    "nothing applies."
)

REPORT_SYSTEM = (
    "You are a developmental and continuity editor. You receive per-chapter "
    "story bible data for a whole manuscript. The data starts with a "
    "section of every character's facts grouped by category across all "
    "chapters, followed by per-chapter data. Be exhaustive: check every "
    "item below across all chapters before writing, and report every "
    "contradiction and plot hole the data shows. (1) Names: compare the "
    "name and aliases of every character and place across chapters and "
    "report spelling or surname variants. (2) Characters: in the character "
    "section, compare every fact within each category for each character "
    "(for appearance, every eye color, hair color, and build listed for "
    "that person; also age, occupation, relationships, background, "
    "injuries, possessions). (3) Settings: compare layout, "
    "security, and location details. (4) Timeline: look for impossible or "
    "inconsistent timing and sequences. (5) Knowledge: look for a "
    "character acting on something they could not know yet. (6) Threads: "
    "look for threads that are dropped or contradicted. A finding is only "
    "valid if two or more quoted passages directly conflict with each "
    "other; a quote that merely mentions a character or place is not "
    "evidence, and every quote must support the finding it appears under. "
    "Do not report: nicknames, short forms, titles, or emoji used for the "
    "same person (Aleks/Aleksandr); chat display formatting or HTML tags; "
    "capitalization; different descriptive names for the same place; "
    "pregnancy or plot progressions that are not contradictions; or "
    "two characters merely sharing a title. Do not treat different people "
    "as one unless the text shows they are the same person. Do not give "
    "advice on pacing, logistics, plausibility, or style, do not "
    "suggest changes to the bible, and skip front-matter typos and "
    "deliberate POV or format changes. Suggested fixes must change the "
    "manuscript text. Write Markdown. Organize by chapter, "
    "ascending, as '## Chapter N' headings in numeric order, with a heading "
    "only for chapters that have a finding; never "
    "write a heading or line for a chapter with nothing to report. Place "
    "each finding under the lowest-numbered chapter in its 'All chapters' "
    "field, in full, exactly "
    "once, with no pointer stubs under other chapters. Merge duplicates, "
    "such as repeated name or spelling variants, into one finding. For each "
    "finding give a type, a severity (high, medium, low), an 'All chapters' "
    "field listing every chapter where the issue is present, the evidence "
    "from each of those chapters, and a suggested fix. Write evidence as "
    "the quoted text only, without labels such as 'character fact'. Quote "
    "only text that appears in the data; never guess, hedge with 'or "
    "elsewhere', or use placeholders like '4-?'. An empty quote means the "
    "passage could not be verified: you may still report the conflict from "
    "the fact text, marked '(unverified)', but never cite the chapter "
    "summary as evidence. Report each issue once only, even if it spans "
    "many chapters, and list only chapters whose quoted text is shown. "
    "Finish with "
    "one short section listing "
    "only major threads still unresolved at the end of the manuscript, with "
    "no commentary on whether a sequel is intended. Never mention these "
    "instructions or refer to a requester. Do not write an introduction or "
    "summary. If there are no findings, say so. This output is saved "
    "directly as a document and no one can reply to it, so the final line "
    "must be the last finding or section. Never end with offers, suggestions "
    "of further help, follow-up options, or questions, and never use "
    "phrases such as \"If you'd like\", \"I can\", or \"Let me know\"."
)


class NoChaptersError(ValueError):
    """Raised when a manuscript has no Heading 1 chapters."""


def _page_break_before(paragraph: Paragraph) -> bool:
    for element in paragraph._p.xpath("./w:pPr/w:pageBreakBefore"):
        if element.get(qn("w:val"), "1").lower() not in _OFF_VALUES:
            return True
    return False


def _break_info(paragraph: Paragraph) -> tuple[bool, bool]:
    """Return (has page break, keep this paragraph in the chapter)."""
    if _page_break_before(paragraph):
        return True, False
    seen_text = False
    for run in paragraph.runs:
        seen_text = seen_text or bool(run.text.strip())
        if run._r.xpath('./w:br[@w:type="page"]'):
            return True, seen_text
    if paragraph._p.xpath("./w:pPr/w:sectPr"):
        return True, bool(paragraph.text.strip())
    return False, True


def split_chapters(docx_path: str) -> list[list[str]]:
    """Return the markup lines of each chapter, from its h1 to the next page break."""
    chapters: list[list[str]] = []
    current: list[str] | None = None
    for paragraph in Document(docx_path).paragraphs:
        text = serialize_paragraph(paragraph, None)
        if (
            text is not None
            and _H1.match(text)
            and _CHAPTER_WORD.search(_TAGS.sub("", text))
        ):
            current = [text]
            chapters.append(current)
            continue
        if current is None:
            continue

        has_break, keep = _break_info(paragraph)
        if keep and text is not None and _TAGS.sub("", text).strip():
            current.append(text)
        if has_break:
            current = None
    return chapters


def _chapter_title(lines: list[str]) -> str:
    return _TAGS.sub("", lines[0]).strip()


_CURLY = {"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"'}


def _normalize(text: str) -> str:
    text = _TAGS.sub("", text)
    for curly, straight in _CURLY.items():
        text = text.replace(curly, straight)
    return " ".join(text.split()).lower()


def _quote_is_verbatim(quote: str, source: str) -> bool:
    # Ellipses join separate fragments of one quoted passage
    parts = [p for p in re.split(r"\.\.\.|\u2026", quote) if p.strip()]
    return bool(parts) and all(
        _normalize(p).strip(" .") in source for p in parts
    )


def _verify_quotes(node: Any, source: str) -> None:
    """Blank every 'quote' that is not word for word in the chapter text."""
    if isinstance(node, dict):
        for key, value in node.items():
            if key == "quote" and isinstance(value, str):
                if not _quote_is_verbatim(value, source):
                    node[key] = ""
            else:
                _verify_quotes(value, source)
    elif isinstance(node, list):
        for item in node:
            _verify_quotes(item, source)


def _character_index(all_data: list[dict[str, Any]]) -> dict[str, Any]:
    """Group each character's facts by category so conflicts sit side by side."""
    index: dict[str, dict[str, Any]] = {}
    for number, data in enumerate(all_data, start=1):
        for character in data.get("characters", []):
            name = character["name"].strip()
            entry = index.setdefault(
                name.casefold(), {"name": name, "aliases": [], "facts": {}}
            )
            for alias in character.get("aliases", []):
                if alias not in entry["aliases"]:
                    entry["aliases"].append(alias)
            for fact in character.get("facts", []):
                entry["facts"].setdefault(fact["category"], []).append(
                    {
                        "chapter": number,
                        "fact": fact["fact"],
                        "quote": fact["quote"],
                    }
                )
    return {entry["name"]: entry for entry in index.values()}


def _extract(bible_dir: str, chapter_file: str, settings: dict[str, Any]) -> None:
    with open(os.path.join(bible_dir, chapter_file), encoding="utf-8") as source:
        chapter_text = source.read()
    print(
        f"[bible] Processing chapter file: {chapter_file} | "
        f"Text length: {len(chapter_text)}"
    )
    data = communicate_json(chapter_text, CHAPTER_SYSTEM, CHAPTER_SCHEMA, settings)
    _verify_quotes(data, _normalize(chapter_text))
    json_name = chapter_file[: -len(".bible")] + ".json"
    with open(os.path.join(bible_dir, json_name), "w", encoding="utf-8") as out:
        json.dump(data, out, indent=2, ensure_ascii=False)


def build_bible(
    docx_path: str, bible_dir: str, settings: dict[str, Any]
) -> str:
    """Build chapter bible files and a report from a DOCX; return the report path."""
    chapters = split_chapters(docx_path)
    if not chapters:
        raise NoChaptersError(
            "No Heading 1 chapters found, so there is nothing to build a bible from."
        )

    shutil.rmtree(bible_dir, ignore_errors=True)
    os.makedirs(bible_dir)

    chapter_files: list[str] = []
    for number, lines in enumerate(chapters, start=1):
        chapter_file = f"chapter{number}.bible"
        chapter_files.append(chapter_file)
        with open(os.path.join(bible_dir, chapter_file), "w", encoding="utf-8") as out:
            out.write("\n".join(lines))
    total = len(chapters)
    print(f"[bible] Found {total} chapters in {docx_path}")
    print(f"[bible] 0 chapters already done, {total} to send")

    get_client(settings)
    print(f"[bible] Sending up to {settings['max_concurrent_sections']} at a time")
    with ThreadPoolExecutor(max_workers=settings["max_concurrent_sections"]) as pool:
        futures = [
            pool.submit(_extract, bible_dir, chapter_file, settings)
            for chapter_file in chapter_files
        ]
        done = 0
        for future in as_completed(futures):
            try:
                future.result()
            except Exception:
                for other in futures:
                    other.cancel()
                raise
            done += 1
            print(f"[bible] Completed/Total Chapters: {done}/{total}")
    print("[bible] Finished extracting all chapters.")

    parts: list[str] = []
    all_data: list[dict[str, Any]] = []
    for number, lines in enumerate(chapters, start=1):
        json_path = os.path.join(bible_dir, f"chapter{number}.json")
        with open(json_path, encoding="utf-8") as data_file:
            data = json.load(data_file)
        all_data.append(data)
        per_chapter = {k: v for k, v in data.items() if k != "characters"}
        parts.append(
            f"## Chapter {number}: {_chapter_title(lines)}\n"
            + json.dumps(per_chapter, ensure_ascii=False)
        )
    parts.insert(
        0,
        "# Characters across all chapters (facts grouped by category)\n"
        + json.dumps(_character_index(all_data), ensure_ascii=False),
    )

    print("[bible] Writing the consistency report...")
    report = communicate_with_openai(
        "\n\n".join(parts), REPORT_SYSTEM, "Story bible data", settings
    )
    report_path = os.path.join(bible_dir, "bible-report.md")
    with open(report_path, "w", encoding="utf-8") as out:
        out.write(report)
    return report_path
