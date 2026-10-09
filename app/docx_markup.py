import hashlib
import html
import os
import re

from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Inches, RGBColor
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.text.run import Run

_IMAGE_NAME = r"[A-Za-z0-9_-][\w.-]*"
# Optional " WIDTHxHEIGHT" in inches
_IMAGE_DIMS = r"(?: (\d+(?:\.\d+)?)x(\d+(?:\.\d+)?))?"
_INLINE_TOKEN = re.compile(
    rf'(<a href="[^"]*">|</a>|</?b>|</?i>|<image {_IMAGE_NAME} inline{_IMAGE_DIMS}>)'
)
_IMAGE_INLINE = re.compile(rf"^<image ({_IMAGE_NAME}) inline{_IMAGE_DIMS}>$")
_IMAGE_LINE = re.compile(
    rf"^<image ({_IMAGE_NAME}) (left|center|right){_IMAGE_DIMS}>$"
)
_EMU_PER_INCH = 914400
_MAX_IMAGE_WIDTH = Inches(6)
_LINK_OPEN = re.compile(r'^<a href="([^"]*)">$')
_HEADING = re.compile(r"^<h([1-9])>(.*)</h\1>$")
_LINK_COLOR = RGBColor(0x05, 0x63, 0xC1)


def _save_run_images(run: Run, image_dir: str) -> str:
    """Save pictures in a run to image_dir and return their inline tags."""
    tags = []
    for drawing in run._r.xpath(".//w:drawing"):
        dims = ""
        extents = drawing.xpath(".//wp:extent")
        if extents:
            width = int(extents[0].get("cx")) / _EMU_PER_INCH
            height = int(extents[0].get("cy")) / _EMU_PER_INCH
            dims = f" {width:.2f}x{height:.2f}"
        for rel_id in drawing.xpath(".//a:blip/@r:embed"):
            image_part = run.part.related_parts[rel_id]
            extension = os.path.splitext(image_part.partname)[1] or ".bin"
            digest = hashlib.sha256(image_part.blob).hexdigest()[:12]
            name = f"img-{digest}{extension}"
            os.makedirs(image_dir, exist_ok=True)
            with open(os.path.join(image_dir, name), "wb") as image_file:
                image_file.write(image_part.blob)
            tags.append(f"<image {name} inline{dims}>")
    return "".join(tags)


def _serialize_run(run: Run, image_dir: str | None = None) -> str:
    text = run.text
    if run.bold and run.italic:
        text = f"<b><i>{text}</i></b>"
    elif run.bold:
        text = f"<b>{text}</b>"
    elif run.italic:
        text = f"<i>{text}</i>"
    if image_dir is not None:
        text += _save_run_images(run, image_dir)
    return text


def _serialize_inline(paragraph: Paragraph, image_dir: str | None = None) -> str:
    parts = []
    for item in paragraph.iter_inner_content():
        if isinstance(item, Hyperlink):
            inner = "".join(_serialize_run(run, image_dir) for run in item.runs)
            if item.url:
                href = html.escape(item.url, quote=True)
                inner = f'<a href="{href}">{inner}</a>'
            parts.append(inner)
        else:
            parts.append(_serialize_run(item, image_dir))
    return "".join(parts)


def _justification(paragraph: Paragraph) -> str | None:
    """Return the raw w:jc value, falling back to the paragraph style chain."""
    # Raw XML: python-docx's alignment enum rejects values such as "start".
    properties = paragraph._p.pPr
    element = properties.find(qn("w:jc")) if properties is not None else None
    style = paragraph.style
    while element is None and style is not None:
        style_properties = style.element.pPr
        if style_properties is not None:
            element = style_properties.find(qn("w:jc"))
        style = style.base_style
    return element.get(qn("w:val")) if element is not None else None


def serialize_paragraph(
    paragraph: Paragraph, image_dir: str | None = None
) -> str | None:
    """Return the markup line for a paragraph, or None if it has no content.

    When image_dir is given, pictures are saved there and emitted as tags.
    """
    if not paragraph.runs and not paragraph.hyperlinks:
        return None

    text = _serialize_inline(paragraph, image_dir)
    image_only = _IMAGE_INLINE.match(text.strip())
    if image_only:
        position = {"center": "center", "right": "right", "end": "right"}.get(
            _justification(paragraph), "left"
        )
        dims = ""
        if image_only.group(2):
            dims = f" {image_only.group(2)}x{image_only.group(3)}"
        return f"<image {image_only.group(1)} {position}{dims}>"
    style_name = paragraph.style.name if paragraph.style else ""
    if style_name == "Title":
        return f"<title>{text}</title>"
    if style_name.startswith("Heading"):
        level = re.search(r"\d+", style_name)
        if level:
            return f"<h{level.group()}>{text}</h{level.group()}>"
    if _justification(paragraph) == "center":
        return f"<center>{text}</center>"
    return f"<p>{text}</p>"


def _curly_quotes(text: str, in_quote: bool) -> tuple[str, bool]:
    out = []
    for char in text:
        if char == '"':
            out.append("\u201d" if in_quote else "\u201c")
            in_quote = not in_quote
        else:
            out.append(char)
    return "".join(out), in_quote


def _new_hyperlink(paragraph: Paragraph, url: str) -> BaseOxmlElement:
    rel_id = paragraph.part.relate_to(
        url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    paragraph._p.append(link)
    return link


def _add_picture(
    paragraph: Paragraph,
    name: str,
    image_dir: str | None,
    dims: tuple[str | None, str | None] = (None, None),
) -> None:
    path = os.path.join(image_dir or "", name)
    if not image_dir or not os.path.isfile(path):
        raise ValueError(f"Image file not found for tag: {name!r}")
    picture = paragraph.add_run().add_picture(path)
    if dims[0] and dims[1]:
        picture.width = Inches(float(dims[0]))
        picture.height = Inches(float(dims[1]))
    elif picture.width > _MAX_IMAGE_WIDTH:
        picture.height = int(picture.height * _MAX_IMAGE_WIDTH / picture.width)
        picture.width = _MAX_IMAGE_WIDTH


def _is_image_only(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing")) and not paragraph.text.strip()


def _add_chapter_page_break(doc: DocumentObject) -> None:
    """Add a page break, placing it before a directly preceding image."""
    previous = doc.paragraphs[-1] if doc.paragraphs else None
    page_break = doc.add_page_break()
    if previous is not None and _is_image_only(previous):
        previous._p.addprevious(page_break._p)


def add_inline_markup(
    paragraph: Paragraph, content: str, image_dir: str | None = None
) -> None:
    """Append runs and hyperlinks described by inline markup to a paragraph."""
    bold = italic = False
    link = None
    in_quote = False
    for token in _INLINE_TOKEN.split(content):
        if not token:
            continue
        link_open = _LINK_OPEN.match(token)
        image = _IMAGE_INLINE.match(token)
        if image:
            _add_picture(
                paragraph, image.group(1), image_dir, (image.group(2), image.group(3))
            )
        elif token == "<b>":
            bold = True
        elif token == "</b>":
            bold = False
        elif token == "<i>":
            italic = True
        elif token == "</i>":
            italic = False
        elif link_open:
            url = html.unescape(link_open.group(1))
            link = _new_hyperlink(paragraph, url) if url else None
        elif token == "</a>":
            link = None
        else:
            text, in_quote = _curly_quotes(token, in_quote)
            run = paragraph.add_run(text)
            run.bold = bold or None
            run.italic = italic or None
            if link is not None:
                run.font.underline = True
                run.font.color.rgb = _LINK_COLOR
                link.append(run._r)


def add_markup_line(
    doc: DocumentObject, line: str, image_dir: str | None = None
) -> bool:
    """Add one markup line to the document; return False if unrecognized."""
    heading = _HEADING.match(line)
    image_line = _IMAGE_LINE.match(line)
    if image_line:
        paragraph = doc.add_paragraph()
        paragraph.alignment = {
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
        }.get(image_line.group(2))
        _add_picture(
            paragraph,
            image_line.group(1),
            image_dir,
            (image_line.group(3), image_line.group(4)),
        )
    elif line.startswith("<title>") and line.endswith("</title>"):
        paragraph = doc.add_heading("", level=1)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline_markup(paragraph, line[7:-8], image_dir)
    elif heading:
        level = int(heading.group(1))
        if level == 1:
            _add_chapter_page_break(doc)
        paragraph = doc.add_heading("", level=level)
        paragraph.paragraph_format.space_after = Inches(
            0.5 - 0.05 * (level - 1)
        )
        if level == 1:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline_markup(paragraph, heading.group(2), image_dir)
    elif line.startswith("<center>") and line.endswith("</center>"):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline_markup(paragraph, line[8:-9], image_dir)
    elif line.startswith("<p>") and line.endswith("</p>"):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Inches(0.20)
        add_inline_markup(paragraph, line[3:-4], image_dir)
    else:
        return False
    return True
