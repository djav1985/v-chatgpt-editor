import os
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from io import StringIO
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_BREAK

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))

from bible import NoChaptersError, build_bible, split_chapters  # noqa: E402

SETTINGS = {"model": "test-model", "max_concurrent_sections": 2}


def make_docx(directory: str, build) -> str:
    doc = Document()
    build(doc)
    path = os.path.join(directory, "book.docx")
    doc.save(path)
    return path


class SplitChaptersTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def split(self, build) -> list[list[str]]:
        return split_chapters(make_docx(self.tmp.name, build))

    def test_chapter_ends_at_page_break_and_excludes_text_after(self) -> None:
        def build(doc) -> None:
            doc.add_heading("Chapter One", 1)
            doc.add_paragraph("Inside.")
            doc.add_page_break()
            doc.add_paragraph("After the break.")
            doc.add_heading("Chapter Two", 1)
            doc.add_paragraph("Second.")

        chapters = self.split(build)

        self.assertEqual(chapters[0], ["<h1>Chapter One</h1>", "<p>Inside.</p>"])
        self.assertEqual(chapters[1], ["<h1>Chapter Two</h1>", "<p>Second.</p>"])

    def test_last_chapter_without_break_runs_to_end(self) -> None:
        def build(doc) -> None:
            doc.add_heading("Chapter One", 1)
            doc.add_paragraph("A.")
            doc.add_heading("Chapter Two", 1)
            doc.add_paragraph("B.")
            doc.add_paragraph("C.")

        chapters = self.split(build)

        self.assertEqual(chapters[0], ["<h1>Chapter One</h1>", "<p>A.</p>"])
        self.assertEqual(
            chapters[1], ["<h1>Chapter Two</h1>", "<p>B.</p>", "<p>C.</p>"]
        )

    def test_text_before_first_h1_is_skipped(self) -> None:
        def build(doc) -> None:
            doc.add_paragraph("Dedication.")
            doc.add_heading("Chapter One", 1)
            doc.add_paragraph("A.")

        self.assertEqual(
            self.split(build), [["<h1>Chapter One</h1>", "<p>A.</p>"]]
        )

    def test_h2_stays_inside_chapter(self) -> None:
        def build(doc) -> None:
            doc.add_heading("Chapter One", 1)
            doc.add_heading("Part", 2)
            doc.add_paragraph("A.")

        self.assertEqual(
            self.split(build),
            [["<h1>Chapter One</h1>", "<h2>Part</h2>", "<p>A.</p>"]],
        )

    def test_break_before_h1_does_not_end_its_chapter(self) -> None:
        def build(doc) -> None:
            doc.add_heading("Chapter One", 1)
            doc.add_paragraph("A.")
            doc.add_page_break()
            doc.add_heading("Chapter Two", 1)
            doc.add_paragraph("B.")

        chapters = self.split(build)

        self.assertEqual(chapters[1], ["<h1>Chapter Two</h1>", "<p>B.</p>"])

    def test_break_at_end_of_paragraph_keeps_its_text(self) -> None:
        def build(doc) -> None:
            doc.add_heading("Chapter One", 1)
            doc.add_paragraph("A.")
            tail = doc.add_paragraph("Last.")
            tail.add_run().add_break(WD_BREAK.PAGE)
            doc.add_paragraph("Excluded.")

        self.assertEqual(
            self.split(build),
            [["<h1>Chapter One</h1>", "<p>A.</p>", "<p>Last.</p>"]],
        )

    def test_page_break_before_property_ends_chapter(self) -> None:
        def build(doc) -> None:
            doc.add_heading("Chapter One", 1)
            doc.add_paragraph("A.")
            excluded = doc.add_paragraph("Excluded.")
            excluded.paragraph_format.page_break_before = True

        self.assertEqual(
            self.split(build), [["<h1>Chapter One</h1>", "<p>A.</p>"]]
        )

    def test_no_h1_gives_no_chapters(self) -> None:
        def build(doc) -> None:
            doc.add_paragraph("Only text.")

        self.assertEqual(self.split(build), [])


class BuildBibleTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.bible_dir = os.path.join(self.tmp.name, "bible")

    def docx(self) -> str:
        def build(doc) -> None:
            doc.add_heading("Chapter One", 1)
            doc.add_paragraph("A.")
            doc.add_page_break()
            doc.add_heading("Chapter Two", 1)
            doc.add_paragraph("B.")

        return make_docx(self.tmp.name, build)

    def test_writes_chapter_files_json_and_report(self) -> None:
        with (
            patch("bible.get_client"),
            patch("bible.communicate_json", return_value={"summary": "s"}) as extract,
            patch("bible.communicate_with_openai", return_value="REPORT") as report,
            redirect_stdout(StringIO()),
        ):
            report_path = build_bible(self.docx(), self.bible_dir, SETTINGS)

        self.assertEqual(
            sorted(os.listdir(self.bible_dir)),
            [
                "bible-report.md",
                "chapter1.bible",
                "chapter1.json",
                "chapter2.bible",
                "chapter2.json",
            ],
        )
        with open(os.path.join(self.bible_dir, "chapter1.bible"), encoding="utf-8") as f:
            self.assertEqual(f.read(), "<h1>Chapter One</h1>\n<p>A.</p>")
        with open(report_path, encoding="utf-8") as f:
            self.assertEqual(f.read(), "REPORT")
        self.assertEqual(extract.call_count, 2)
        report_input = report.call_args.args[0]
        self.assertIn("## Chapter 1: Chapter One", report_input)
        self.assertIn("## Chapter 2: Chapter Two", report_input)

    def test_rebuilds_directory_each_run(self) -> None:
        os.makedirs(self.bible_dir)
        stale = os.path.join(self.bible_dir, "chapter9.json")
        with open(stale, "w", encoding="utf-8") as f:
            f.write("{}")

        with (
            patch("bible.get_client"),
            patch("bible.communicate_json", return_value={}),
            patch("bible.communicate_with_openai", return_value="R"),
            redirect_stdout(StringIO()),
        ):
            build_bible(self.docx(), self.bible_dir, SETTINGS)

        self.assertFalse(os.path.exists(stale))

    def test_extraction_failure_aborts_without_report(self) -> None:
        with (
            patch("bible.get_client"),
            patch("bible.communicate_json", side_effect=RuntimeError("boom")),
            patch("bible.communicate_with_openai") as report,
            redirect_stdout(StringIO()),
        ):
            with self.assertRaisesRegex(RuntimeError, "boom"):
                build_bible(self.docx(), self.bible_dir, SETTINGS)

        report.assert_not_called()

    def test_no_h1_raises_and_makes_no_calls(self) -> None:
        path = make_docx(self.tmp.name, lambda doc: doc.add_paragraph("Text."))
        with patch("bible.communicate_json") as extract:
            with self.assertRaises(NoChaptersError):
                build_bible(path, self.bible_dir, SETTINGS)

        extract.assert_not_called()
        self.assertFalse(os.path.exists(self.bible_dir))


if __name__ == "__main__":
    unittest.main()
