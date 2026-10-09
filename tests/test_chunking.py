import os
import tempfile
import sys
import unittest
import warnings
from unittest.mock import patch

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))

import tiktoken
from docx import Document

from docx_handler import split_into_sections  # noqa: E402
from token_count import count_model_tokens  # noqa: E402


def count_whitespace_tokens(text: str) -> int:
    """Provide deterministic counts for section-boundary tests."""
    return len(text.split())


class TokenCountingTests(unittest.TestCase):
    def test_counts_using_the_configured_model_encoding(self) -> None:
        text = "Some serialized <p>manuscript text</p>."
        expected = len(tiktoken.encoding_for_model("gpt-5-mini").encode(text))

        self.assertEqual(count_model_tokens(text, "gpt-5-mini"), expected)

    def test_rejects_models_without_a_known_encoding(self) -> None:
        with self.assertRaisesRegex(ValueError, "not configured for model"):
            count_model_tokens("text", "unsupported-model-for-test")


class ParagraphChunkingTests(unittest.TestCase):
    def split_paragraphs(
        self, paragraphs: list[str], max_tokens: int
    ) -> tuple[list[list[str]], list[str], list[str]]:
        with tempfile.TemporaryDirectory() as temp_dir:
            original_directory = os.getcwd()
            try:
                os.chdir(temp_dir)
                document = Document()
                for text in paragraphs:
                    document.add_paragraph(text)
                filename = os.path.join(temp_dir, "fixture.docx")
                document.save(filename)

                counted_texts: list[str] = []

                def count_and_record(text: str) -> int:
                    counted_texts.append(text)
                    return count_whitespace_tokens(text)

                with patch(
                    "docx_handler.count_model_tokens",
                    side_effect=count_and_record,
                ):
                    sections = split_into_sections(filename, max_tokens)

                saved_sections = []
                for index in range(1, len(sections) + 1):
                    section_path = os.path.join(
                        temp_dir, "tmp", "fixture", f"{index}-section.old"
                    )
                    with open(section_path, encoding="utf-8") as section_file:
                        saved_sections.append(section_file.read())
                return sections, saved_sections, counted_texts
            finally:
                os.chdir(original_directory)

    def test_packs_whole_paragraphs_up_to_the_budget(self) -> None:
        sections, saved_sections, counted_texts = self.split_paragraphs(
            ["one two", "three four", "five"], 4
        )

        self.assertEqual(
            sections,
            [
                ["<p>one two</p>", "<p>three four</p>"],
                ["<p>five</p>"],
            ],
        )
        self.assertEqual(
            saved_sections,
            ["<p>one two</p>\n<p>three four</p>", "<p>five</p>"],
        )
        self.assertIn(saved_sections[0], counted_texts)

    def test_keeps_oversized_first_paragraph_without_empty_chunk(self) -> None:
        with warnings.catch_warnings(record=True) as caught_warnings:
            warnings.simplefilter("always")
            sections, saved_sections, _ = self.split_paragraphs(
                ["one two three", "four"], 2
            )

        self.assertEqual(
            sections,
            [["<p>one two three</p>"], ["<p>four</p>"]],
        )
        self.assertEqual(len(saved_sections), 2)
        self.assertTrue(
            any(
                "Paragraph 1 contains 3 tokens" in str(item.message)
                for item in caught_warnings
            )
        )

    def test_warns_for_oversized_later_paragraph_and_preserves_order(self) -> None:
        with warnings.catch_warnings(record=True) as caught_warnings:
            warnings.simplefilter("always")
            sections, _, _ = self.split_paragraphs(["first", "two three four"], 2)

        self.assertEqual(
            sections,
            [["<p>first</p>"], ["<p>two three four</p>"]],
        )
        self.assertTrue(
            any(
                "Paragraph 2 contains 3 tokens" in str(item.message)
                for item in caught_warnings
            )
        )

    def test_empty_document_produces_no_sections(self) -> None:
        sections, saved_sections, _ = self.split_paragraphs([], 10)

        self.assertEqual(sections, [])
        self.assertEqual(saved_sections, [])

    def test_translations_use_separate_workspaces_for_the_same_source(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            original_directory = os.getcwd()
            try:
                os.chdir(temp_dir)
                document = Document()
                document.add_paragraph("shared source")
                filename = "fixture.docx"
                document.save(filename)

                french_dir = os.path.join("tmp", "FRENCH_fixture")
                german_dir = os.path.join("tmp", "GERMAN_fixture")
                split_into_sections(filename, 10, tmp_dir=french_dir)
                split_into_sections(filename, 10, tmp_dir=german_dir)

                self.assertTrue(
                    os.path.exists(os.path.join(french_dir, "1-section.old"))
                )
                self.assertTrue(
                    os.path.exists(os.path.join(german_dir, "1-section.old"))
                )
                self.assertFalse(
                    os.path.exists(os.path.join("tmp", "fixture", "1-section.old"))
                )
            finally:
                os.chdir(original_directory)

    def test_requires_a_positive_budget(self) -> None:
        with self.assertRaisesRegex(ValueError, "positive integer"):
            split_into_sections("unused.docx", 0)

    def test_changed_source_clears_stale_sections_and_results(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            original_directory = os.getcwd()
            try:
                os.chdir(temp_dir)
                filename = "fixture.docx"
                document = Document()
                document.add_paragraph("first")
                document.add_paragraph("second")
                document.save(filename)
                split_into_sections(filename, 10)

                section_dir = os.path.join("tmp", "fixture")
                stale_result = os.path.join(section_dir, "1-section.new")
                with open(stale_result, "w", encoding="utf-8") as result_file:
                    result_file.write("stale result")

                document = Document()
                document.add_paragraph("replacement")
                document.save(filename)
                split_into_sections(filename, 10)

                self.assertFalse(os.path.exists(stale_result))
                self.assertFalse(
                    os.path.exists(os.path.join(section_dir, "2-section.old"))
                )
            finally:
                os.chdir(original_directory)


if __name__ == "__main__":
    unittest.main()
