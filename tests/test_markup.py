import base64
import io
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))

os.environ["OPENAI_API_KEY"] = "test-api-key"

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from docx_handler import merge_groups_and_save  # noqa: E402
from docx_markup import add_markup_line, serialize_paragraph  # noqa: E402


def round_trip(lines: list[str]) -> list[str]:
    doc = Document()
    for line in lines:
        assert add_markup_line(doc, line)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    serialized = [serialize_paragraph(p) for p in Document(buffer).paragraphs]
    return [line for line in serialized if line is not None]


class HyperlinkTests(unittest.TestCase):
    def test_start_alignment_serializes_as_normal_paragraph(self) -> None:
        paragraph = Document().add_paragraph("Start-aligned text")
        justification = OxmlElement("w:jc")
        justification.set(qn("w:val"), "start")
        paragraph._p.get_or_add_pPr().append(justification)

        self.assertEqual(serialize_paragraph(paragraph), "<p>Start-aligned text</p>")

    def test_hyperlink_round_trips(self) -> None:
        line = (
            '<p>See <a href="https://example.com/a?b=1&amp;c=2">the site</a> now</p>'
        )
        self.assertEqual(round_trip([line]), [line])

    def test_hyperlink_with_styled_text_round_trips(self) -> None:
        line = '<p><a href="https://example.com/">a <b><i>bold</i></b> link</a></p>'
        self.assertEqual(round_trip([line]), [line])

    def test_hyperlink_only_paragraph_is_not_empty(self) -> None:
        line = '<p><a href="https://example.com/">only link</a></p>'
        self.assertEqual(round_trip([line]), [line])

    def test_link_in_centered_paragraph(self) -> None:
        line = '<center><a href="https://example.com/">mid</a></center>'
        self.assertEqual(round_trip([line]), [line])

    def test_empty_href_becomes_plain_text(self) -> None:
        self.assertEqual(
            round_trip(['<p><a href="">plain</a></p>']), ["<p>plain</p>"]
        )


class HeadingTests(unittest.TestCase):
    def test_heading_spacing_decreases_with_heading_level(self) -> None:
        for level in range(1, 10):
            with self.subTest(level=level):
                document = Document()
                self.assertTrue(
                    add_markup_line(document, f"<h{level}>Heading</h{level}>")
                )
                heading = next(
                    paragraph
                    for paragraph in document.paragraphs
                    if paragraph.style.name == f"Heading {level}"
                )

                self.assertAlmostEqual(
                    heading.paragraph_format.space_after.inches,
                    0.5 - 0.05 * (level - 1),
                )

    def test_h3_to_h6_round_trip(self) -> None:
        lines = [f"<h{n}>Level {n}</h{n}>" for n in range(3, 7)]
        self.assertEqual(round_trip(lines), lines)

    def test_h2_with_inline_markup(self) -> None:
        line = "<h2>A <i>styled</i> heading</h2>"
        self.assertEqual(round_trip([line]), [line])

    def test_unrecognized_line_reports_false(self) -> None:
        self.assertFalse(add_markup_line(Document(), "<div>x</div>"))


class InlineTests(unittest.TestCase):
    def test_bold_italic_nested_keeps_both(self) -> None:
        line = "<p>x <b><i>both</i></b> y</p>"
        self.assertEqual(round_trip([line]), [line])

    def test_quotes_become_curly_across_inline_tags(self) -> None:
        result = round_trip(['<p>"Hello <i>there</i>"</p>'])
        self.assertEqual(result, ["<p>\u201cHello <i>there</i>\u201d</p>"])


class ImageTests(unittest.TestCase):
    PNG = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )

    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.image_dir = self.temp_dir.name
        self.png_path = os.path.join(self.image_dir, "src.png")
        with open(self.png_path, "wb") as png_file:
            png_file.write(self.PNG)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def build_doc(self) -> Document:
        doc = Document()
        doc.add_heading("Chapter", level=2)
        doc.add_paragraph().add_run().add_picture(self.png_path)
        doc.add_paragraph("After image")
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return Document(buffer)

    def test_image_between_heading_and_paragraph_keeps_order(self) -> None:
        lines = [
            serialize_paragraph(p, self.image_dir)
            for p in self.build_doc().paragraphs
        ]
        self.assertRegex(lines[0], r"^<h2>Chapter</h2>$")
        self.assertRegex(
            lines[1], r"^<image img-[0-9a-f]{12}\.png left \d+\.\d\d x\d+\.\d\d>$".replace(" x", "x")
        )
        self.assertEqual(lines[2], "<p>After image</p>")
        self.assertTrue(
            os.path.exists(os.path.join(self.image_dir, lines[1].split()[1]))
        )

    def test_image_round_trips_through_merge_markup(self) -> None:
        lines = [
            serialize_paragraph(p, self.image_dir)
            for p in self.build_doc().paragraphs
        ]
        rebuilt = Document()
        for line in lines:
            self.assertTrue(add_markup_line(rebuilt, line, self.image_dir))
        again = [serialize_paragraph(p, self.image_dir) for p in rebuilt.paragraphs]
        self.assertEqual(again[1:], lines[1:])
        self.assertEqual(len(rebuilt.inline_shapes), 1)

    def test_inline_image_within_text(self) -> None:
        doc = Document()
        self.assertTrue(add_markup_line(doc, "<p>A</p>", self.image_dir))
        name = "img-x.png"
        os.replace(self.png_path, os.path.join(self.image_dir, name))
        self.assertTrue(
            add_markup_line(doc, f"<p>See <image {name} inline> here</p>", self.image_dir)
        )
        self.assertEqual(len(doc.inline_shapes), 1)

    def test_alignment_variants_serialize_as_positions(self) -> None:
        doc = Document()
        style = doc.styles.add_style("CenterImg", 1)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cases = [
            (None, None, "left"),
            (WD_ALIGN_PARAGRAPH.CENTER, None, "center"),
            (WD_ALIGN_PARAGRAPH.RIGHT, None, "right"),
            (None, "CenterImg", "center"),
        ]
        for alignment, style_name, expected in cases:
            with self.subTest(expected=expected, style=style_name):
                paragraph = doc.add_paragraph(style=style_name)
                paragraph.alignment = alignment
                paragraph.add_run().add_picture(self.png_path)
                self.assertRegex(
                    serialize_paragraph(paragraph, self.image_dir),
                    rf"^<image img-[0-9a-f]{{12}}\.png {expected} ",
                )

    def test_style_with_start_alignment_does_not_raise(self) -> None:
        doc = Document()
        style = doc.styles.add_style("StartStyle", 1)
        justification = OxmlElement("w:jc")
        justification.set(qn("w:val"), "start")
        style.element.get_or_add_pPr().append(justification)
        paragraph = doc.add_paragraph(style="StartStyle")
        paragraph.add_run().add_picture(self.png_path)
        self.assertRegex(
            serialize_paragraph(paragraph, self.image_dir), r"^<image \S+ left "
        )

    def test_image_with_whitespace_run_is_still_standalone(self) -> None:
        paragraph = Document().add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(" ")
        paragraph.add_run().add_picture(self.png_path)
        self.assertRegex(
            serialize_paragraph(paragraph, self.image_dir),
            r"^<image img-[0-9a-f]{12}\.png center ",
        )

    def test_missing_image_file_raises(self) -> None:
        with self.assertRaisesRegex(ValueError, "Image file not found"):
            add_markup_line(Document(), "<image img-nope.png center>", self.image_dir)

    def test_dimensions_are_applied_and_round_trip(self) -> None:
        name = "img-x.png"
        os.replace(self.png_path, os.path.join(self.image_dir, name))
        doc = Document()
        self.assertTrue(
            add_markup_line(doc, f"<image {name} center 3.50x2.25>", self.image_dir)
        )
        shape = doc.inline_shapes[0]
        self.assertAlmostEqual(shape.width.inches, 3.5, places=2)
        self.assertAlmostEqual(shape.height.inches, 2.25, places=2)
        self.assertRegex(
            serialize_paragraph(doc.paragraphs[0], self.image_dir),
            r"^<image img-[0-9a-f]{12}\.png center 3\.50x2\.25>$",
        )

    def test_page_break_moves_before_image_preceding_h1(self) -> None:
        name = "img-x.png"
        os.replace(self.png_path, os.path.join(self.image_dir, name))
        doc = Document()
        for line in ["<p>End</p>", f"<image {name} center>", "<h1>Two</h1>"]:
            self.assertTrue(add_markup_line(doc, line, self.image_dir))
        kinds = []
        for paragraph in doc.paragraphs:
            if paragraph._p.xpath('.//w:br[@w:type="page"]'):
                kinds.append("break")
            elif paragraph._p.xpath(".//w:drawing"):
                kinds.append("image")
            else:
                kinds.append(paragraph.text)
        self.assertEqual(kinds, ["End", "break", "image", "Two"])


class MergeTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_directory = os.getcwd()
        os.chdir(self.temp_dir.name)
        os.makedirs(os.path.join("tmp", "fixture"))

    def tearDown(self) -> None:
        os.chdir(self.original_directory)
        self.temp_dir.cleanup()

    def test_merge_preserves_headings_and_hyperlinks(self) -> None:
        lines = [
            "<h2>Chapter <i>One</i></h2>",
            '<p>See <a href="https://example.com/">the site</a> now</p>',
        ]
        with open(
            os.path.join("tmp", "fixture", "1-section.new"),
            "w",
            encoding="utf-8",
        ) as section_file:
            section_file.write("\n".join(lines))

        merge_groups_and_save("fixture.docx", "EDIT", "./output")

        saved_doc = Document(os.path.join("output", "EDIT_fixture.docx"))
        serialized = [
            serialize_paragraph(paragraph) for paragraph in saved_doc.paragraphs
        ]
        self.assertEqual(serialized, lines)
        self.assertTrue(os.path.exists(os.path.join("output", "ORIGINAL_fixture.docx")))

    def test_translation_merge_does_not_create_original_copy(self) -> None:
        translation_dir = os.path.join("tmp", "FRENCH_fixture")
        os.makedirs(translation_dir)
        with open(
            os.path.join(translation_dir, "1-section.new"),
            "w",
            encoding="utf-8",
        ) as section_file:
            section_file.write("<p>Bonjour</p>")
        with open(
            os.path.join(translation_dir, "1-section.old"),
            "w",
            encoding="utf-8",
        ) as section_file:
            section_file.write("<p>Hello</p>")

        merge_groups_and_save(
            "fixture.docx",
            "French",
            "./output",
            translation_dir,
            include_original=False,
        )

        self.assertTrue(os.path.exists(os.path.join("output", "FRENCH_fixture.docx")))
        self.assertFalse(
            os.path.exists(os.path.join("output", "ORIGINAL_fixture.docx"))
        )

    def test_merge_rejects_unrecognized_markup(self) -> None:
        with open(
            os.path.join("tmp", "fixture", "1-section.new"),
            "w",
            encoding="utf-8",
        ) as section_file:
            section_file.write("<div>unsupported</div>")

        with self.assertRaisesRegex(ValueError, "Unrecognized markup"):
            merge_groups_and_save("fixture.docx", "EDIT", "./output")


if __name__ == "__main__":
    unittest.main()
